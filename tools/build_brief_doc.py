# -*- coding: utf-8 -*-
"""
生成精简版（约 4 页）：《计算机类新生选课与培养方案问答助手 —— 方案摘要》
输出: docs/计算机类新生选课与培养方案问答助手_方案摘要_4页版.docx
内置版面高度估算，用于控制篇幅。
"""
import math
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- 版面参数
PAGE_W, PAGE_H = 21.0, 29.7
MARGIN = 2.0
BODY_PT = 10.0
TABLE_PT = 9.0
LINE_SPACING = 1.28
USABLE_W = PAGE_W - 2 * MARGIN          # cm
USABLE_H = PAGE_H - 2 * MARGIN          # cm
CH_CM = lambda pt: pt * 0.03528         # 一个全角字宽 cm
# 中文正文字体单倍行高约为 1.3 em，再乘以行距倍数
LINE_CM = lambda pt, ls: pt * 1.3 * ls / 28.35

EST = {"cm": 0.0}                        # 版面高度累计


def _est_text(text, pt=BODY_PT, indent_cm=0.0, ls=LINE_SPACING, space_after_cm=0.0):
    """估算一段文字的排版高度（cm）。"""
    usable = USABLE_W - indent_cm
    per_line = max(1.0, usable / CH_CM(pt))
    # 全角计 1，半角计 0.5
    units = sum(1.0 if ord(c) > 0x2E7F else 0.5 for c in text)
    lines = max(1, math.ceil(units / per_line))
    EST["cm"] += lines * LINE_CM(pt, ls) + space_after_cm


def _est_row(cells, pt=TABLE_PT, widths=None):
    """估算一个表格行的高度（cm）。"""
    usable = USABLE_W
    max_lines = 1
    for i, txt in enumerate(cells):
        w = widths[i] if widths else usable / len(cells)
        per_line = max(1.0, w / CH_CM(pt))
        units = sum(1.0 if ord(c) > 0x2E7F else 0.5 for c in str(txt))
        max_lines = max(max_lines, math.ceil(units / per_line))
    EST["cm"] += max_lines * LINE_CM(pt, 1.05) + 0.12


# ---------------------------------------------------------------- 样式工具
BODY_EA, BODY_LATIN = "宋体", "Times New Roman"
HEAD_EA, HEAD_LATIN = "黑体", "Arial"
MONO = "Consolas"
NAVY = RGBColor(0x17, 0x37, 0x5E)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(PAGE_W), Cm(PAGE_H)
sec.top_margin = sec.bottom_margin = Cm(MARGIN)
sec.left_margin = sec.right_margin = Cm(MARGIN)


def _set_font(run, ea=BODY_EA, latin=BODY_LATIN, size=BODY_PT, bold=False,
              color=DARK):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), ea)


def _rich(paragraph, text, ea=BODY_EA, latin=BODY_LATIN, size=BODY_PT, color=DARK):
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`|<br\s*/?>)", text):
        if not tok:
            continue
        if re.fullmatch(r"<br\s*/?>", tok):
            r = paragraph.add_run()
            _set_font(r, ea=ea, latin=latin, size=size, color=color)
            r.add_break()
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            _set_font(paragraph.add_run(tok[2:-2]), ea=ea, latin=latin,
                      size=size, bold=True, color=color)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            _set_font(paragraph.add_run(tok[1:-1]), ea=MONO, latin=MONO,
                      size=size - 0.5, color=ACCENT)
        else:
            _set_font(paragraph.add_run(tok), ea=ea, latin=latin, size=size, color=color)
    return paragraph


def h(text, level=1):
    p = doc.add_paragraph(style="Heading %d" % level)
    size = 13.5 if level == 1 else 11.0
    _set_font(p.add_run(text), ea=HEAD_EA, latin=HEAD_LATIN, size=size,
              bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4 if level == 1 else 3)
    p.paragraph_format.line_spacing = 1.15
    EST["cm"] += (size * 1.2 * 1.15 / 28.35) + (9 if level == 1 else 6) / 28.35 + (4 if level == 1 else 3) / 28.35
    return p


def para(text, size=BODY_PT, indent=0.0, space_after=4, align=None, color=DARK):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(space_after)
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    _rich(p, text, size=size, color=color)
    _est_text(re.sub(r"\*\*|`", "", text), size, indent, LINE_SPACING, space_after / 28.35)
    return p


def bullet(text, size=BODY_PT):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(2)
    pf.left_indent = Cm(0.55)
    _rich(p, text, size=size)
    _est_text("• " + re.sub(r"\*\*|`", "", text), size, 0.75, LINE_SPACING, 2 / 28.35)
    return p


def numbered(text, size=BODY_PT):
    p = doc.add_paragraph(style="List Number")
    pf = p.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(2)
    pf.left_indent = Cm(0.55)
    _rich(p, text, size=size)
    _est_text("0. " + re.sub(r"\*\*|`", "", text), size, 0.75, LINE_SPACING, 2 / 28.35)
    return p


def shade(obj, hexcolor):
    el = getattr(obj, "_element", obj)          # 兼容 _Cell / _Paragraph / XML 元素
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    pr.append(shd)


def table(headers, rows, widths, size=TABLE_PT):
    # 列宽按可用宽度归一化，与页边距解耦
    scale = USABLE_W / float(sum(widths))
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        _set_font(p.add_run(htxt), ea=HEAD_EA, latin=HEAD_LATIN, size=size,
                  bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, "17375E")
    _est_row(headers, size, widths)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            _rich(p, str(val), size=size)
        if ri % 2 == 1:
            for c in cells:
                shade(c, "F2F5F9")
        _est_row(row, size, widths)
    for ci, w in enumerate(widths):
        for row in t.rows:
            row.cells[ci].width = Cm(w)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    _set_font(sp.add_run(""), size=3)          # 表格后的微小间隔
    EST["cm"] += 0.18
    return t


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    _set_font(p.add_run(text), ea=BODY_EA, size=8, color=GREY)
    EST["cm"] += 8 * 1.2 / 28.35 + 7 / 28.35
    return p


def layer(title_text, body_text, fill="EAF2FB"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    _set_font(p.add_run(title_text), ea=HEAD_EA, latin=HEAD_LATIN, size=9.5,
              bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    p2.paragraph_format.line_spacing = 1.15
    _rich(p2, body_text, size=8.5)
    _est_row([title_text + " / " + body_text], 8.5, [USABLE_W])
    EST["cm"] += 0.08
    return t


def flow(text):
    """一条横向流程说明条。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "F5F8FC")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    _rich(p, text, size=8.5)
    _est_row([re.sub(r"\*\*|`", "", text)], 8.5, [USABLE_W])
    EST["cm"] += 0.1
    return t


def code_block(text, size=8.0):
    lines = text.strip("\n").split("\n")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "F7F9FC")
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _set_font(p.add_run(line if line else " "), ea=MONO, latin=MONO,
                  size=size, color=RGBColor(0x1B, 0x2A, 0x3A))
    EST["cm"] += (len(lines) * size * 1.2 / 28.35) + 0.14
    return t


# 默认样式
normal = doc.styles["Normal"]
normal.font.name = BODY_LATIN
normal.font.size = Pt(BODY_PT)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_EA)
for lvl, sz in ((1, 13.5), (2, 11.0)):
    st = doc.styles["Heading %d" % lvl]
    st.font.name = HEAD_LATIN
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = NAVY
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), HEAD_EA)

# 页脚页码
fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set_font(fp.add_run("— "), size=8, color=GREY)
_r = fp.add_run()
_set_font(_r, size=8, color=GREY)
for e in (OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")):
    _r._element.append(e)
_r._element[0].set(qn("w:fldCharType"), "begin")
_r._element[1].set(qn("xml:space"), "preserve")
_r._element[1].text = "PAGE"
_r._element[2].set(qn("w:fldCharType"), "end")
_set_font(fp.add_run(" —"), size=8, color=GREY)

# ================================================================ 标题区
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
_set_font(p.add_run("计算机类新生选课与培养方案问答助手"), ea=HEAD_EA,
          latin=HEAD_LATIN, size=17, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
_set_font(p.add_run("基于大模型与检索增强生成（RAG）· 方案摘要（精简版）"), ea=HEAD_EA,
          latin=HEAD_LATIN, size=11, bold=True, color=ACCENT)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
_set_font(p.add_run("V0.1 草案　·　%s　·　交付形态：DSH 插件 + 专属 Skill + 领域知识库"
                    % datetime.now().strftime("%Y 年 %m 月")),
          ea=BODY_EA, size=8.5, color=GREY)
EST["cm"] += 17 * 1.2 / 28.35 + 11 * 1.2 / 28.35 + 8.5 * 1.2 / 28.35 + 0.35

para("**一句话概述**：把培养方案与教师数据结构化入库并向量化，由大模型 API 负责理解与表达，"
     "由专属 Skill 负责领域流程，最终封装为 DSH 插件交付使用。", size=9.5, space_after=6)

# ================================================================ 一
h("一、项目概述", 1)

h("1.1 问题", 2)
bullet("**培养方案难读**：几十页文本、术语密集（必修／限选／模块／方向课组／先修），新生读完仍不知道本学期该选什么。")
bullet("**规则分散且嵌套**：教务系统规则、学院规定、专业要求、开课计划散落在不同文件，彼此还有优先级关系。")
bullet("**信息靠口口相传**：辅导员、学长、群聊口径不一、时效滞后，且没有一条能追到原始条款；教师与课程情况靠私下打听，样本小、易失真。")
bullet("**规划缺失**：先修顺序错误、学分结构不达标、大四被迫补课年年重复；选课周重复答疑占用辅导员大量时间。")

h("1.2 目标与边界", 2)
bullet("**目标**：用自然语言提问即可获得准确、可溯源、个性化的回答，覆盖培养方案解读、课程查询、教师与课程客观信息、规则解释、学分核算、选课规划。")
bullet("**可溯源**：每条事实性结论都能指回原文的章节或条款；检索不到依据就明说，不编造。")
bullet("**交付**：以 DSH 插件形态提供，复用其会话、工具调用、Skill 加载与界面能力，装即用、可版本化分发。")
bullet("**边界**：不做抢课代选、不做教师优劣排名、不做成绩预测；答复以教务系统与学院正式文件为准。")

# ================================================================ 二
h("二、总体方案", 1)

h("2.1 三方分工", 2)
table(
    ["组成", "职责", "关键点"],
    [
        ["关系数据库", "提供确定的事实", "学分、学时、开课学期、先修关系、教师授课记录一律查表，不由模型生成"],
        ["向量知识库 + RAG", "提供有依据的解释", "方案正文、管理规定、教学大纲分块向量化；混合检索 + 重排"],
        ["大模型 API", "理解与表达", "不负责“记住”培养方案，只负责理解问题、组织语言、标注引用"],
        ["专属 Skill", "领域流程", "规定查哪张表、检索什么、如何输出、什么情况下必须拒答"],
    ],
    widths=[2.9, 3.2, 10.5],
)
caption("表 1　三方分工")

h("2.2 分层架构", 2)
layer("① 交互层　DSH Web GUI / CLI",
      "自然语言提问　·　快捷指令　·　课程卡 / 教师卡 / 来源引用　·　追问建议与反馈")
layer("② 插件层　DSH Plugin: class-choice-assistant",
      "工具注册　·　命令注册　·　Skill 目录　·　插件配置　·　会话上下文与学生画像")
layer("③ 技能层　Skills（领域流程编排）",
      "培养方案解读 · 课程查询 · 教师画像 · 学分核算 · 选课规划 · 规则解释 · 通知问答 · 转人工指引")
layer("④ 检索层　RAG Pipeline",
      "查询理解与路由 → 结构化 / 向量 / 混合检索 → 重排与去重 → 上下文组装（带引用元数据）→ 结果校验")
layer("⑤ 模型层与数据层",
      "LLM API（理解与生成）· Embedding API（向量化）· Rerank（可选）　｜　"
      "关系库（方案/课程/教师/规则）· 向量库 · 原始文件库 · 日志与反馈库", fill="EAF7EE")

h("2.3 一次问答的链路", 2)
flow("提问 → 注入学生画像 → 意图识别 → 路由（查表 / 检索 / 混合）→ 重排去重 → 组装提示词 → 模型生成 → "
     "**输出校验** → 返回答案 + 来源卡片 → 日志留痕")
para("**关键差异点**：“输出校验”一步会程序化复核引用编号的真实性与学分等数字的一致性，不一致就重新生成或降级为"
     "仅返回原始数据。教务场景里，**一个错误的学分数字比一句“我不确定”危害大得多**。", size=9.5)

# ================================================================ 三
h("三、数据层", 1)

h("3.1 数据来源", 2)
table(
    ["类别", "内容", "更新频率", "优先级"],
    [
        ["培养方案", "方案全文，含学分结构、课程设置表、修读说明", "每年", "核心"],
        ["课程库", "课程代码、名称、学分、学时、性质、建议学期", "每学期", "核心"],
        ["方案-课程关联", "模块归属、必修／限选／任选、建议修读学期", "每年", "核心"],
        ["先修关系", "课程之间的先修、并修、互斥关系", "每年", "核心"],
        ["教师信息与授课记录", "公开信息 + 历年授课课程与学期", "每学年／每学期", "核心"],
        ["教学大纲", "课程目标、内容纲要、考核方式与成绩构成", "不定期", "高"],
        ["开课计划与课表", "时间、地点、容量、任课教师", "每学期", "高"],
        ["规则与通知", "学分上限、重修免修、学业预警、学院通知", "每年／实时", "高"],
    ],
    widths=[3.0, 8.4, 3.0, 2.2],
)
caption("表 2　数据来源清单（示例）")

h("3.2 数据处理流水线", 2)
flow("采集 → 解析（PDF／Word／Excel 转文本 + 表格还原）→ 清洗（去页眉页脚、合并跨页表）→ "
     "**LLM 辅助结构化抽取 + 关键字段 100% 人工复核** → 入库 → 语义分块与向量化 → 版本冻结与灰度发布")
para("风险最高的环节是结构化抽取：**学分、学时、课程性质一旦抽错，后续所有学分核算都是错的**，"
     "因此关键字段必须逐一复核，而不能只做抽样。", size=9.5)

h("3.3 核心表与知识库", 2)
para("核心表包括专业与方案版本、课程主数据、方案-课程关联、模块学分要求、先修关系、开课记录、教师与授课历史、"
     "考核方式、规则条款、文档分块以及问答日志。", size=9.5)
bullet("**分块**：按章节／条款语义切分，单块 400～800 字、重叠 10%～15%，每块携带文档名、版本、专业、年级、章节路径与页码。", size=9.5)
bullet("**检索**：向量检索与关键词倒排混合，支持按专业、年级、方案版本做元数据过滤。", size=9.5)
bullet("**版本**：以“专业 + 适用年级 + 生效学年”作为版本标识，默认只在学生适用版本内检索，避免用新方案回答老生问题。", size=9.5)

# ================================================================ 四
h("四、检索增强生成（RAG）", 1)

h("4.1 意图路由", 2)
table(
    ["意图类型", "示例问题", "检索通道"],
    [
        ["事实查询", "数据结构这门课多少学分？", "关系库精确查询"],
        ["关系查询", "学人工智能导论之前要先修什么？", "先修关系表"],
        ["规则解释", "培养方案里的“专业方向课组”是什么意思？", "条款库 + 文档检索"],
        ["核算类", "我现在还差多少学分才能毕业？", "关系库计算（需个人学业数据授权）"],
        ["冲突判断", "选这门课会不会和必修课撞时间？", "开课时间比对"],
        ["越界问题", "你觉得这门课水不水？", "不进入检索，按边界话术回应"],
    ],
    widths=[2.4, 7.6, 6.6],
)
caption("表 3　意图路由策略")
para("**准确率提升最直接的手段是课程别名表**：把“程设”“程序设计基础”“C 语言”等不同叫法归一到同一门课程，"
     "否则新生用口语提问时很难召回正确结果。", size=9.5)

h("4.2 生成约束", 2)
bullet("**只依据证据包作答**：证据不足时明确说明“现有资料中未找到相关说明”，检索分数低于阈值时直接走拒答与转人工，不兜底编造。", size=9.5)
bullet("**引用强制**：每条事实性结论标注引用编号，生成后程序化校验编号真实存在，无效引用触发重新生成。", size=9.5)
bullet("**数字复核**：学分、学时、门数等必须与查表结果逐字一致，不一致则替换为重查结果。", size=9.5)
bullet("**表达风格**：面向新生，先给结论再给依据，少用术语，长答案分点呈现，并注明所依据的方案版本。", size=9.5)

# ================================================================ 五
h("五、专属 Skill 与 DSH 插件", 1)

h("5.1 Skill 清单与定位", 2)
table(
    ["Skill", "解决问题"],
    [
        ["培养方案解读 / 课程信息查询", "方案怎么读、学分结构是什么；某门课多少学分、何时开、谁在上"],
        ["教师与课程画像", "教师公开信息、研究方向、授课历史、课程考核方式"],
        ["学分与毕业要求核算", "已修多少、还差多少、哪个模块不达标"],
        ["选课规划建议", "下学期建议选什么、先修顺序怎么排、有无冲突风险"],
        ["规则条款解释 / 通知问答", "选课规则、重修免修怎么规定；最新通知说了什么"],
        ["转人工与办事指引", "助手答不了的事该找哪个部门、在哪里办、需要什么材料"],
    ],
    widths=[5.0, 11.6],
)
caption("表 4　专属 Skill 清单")
para("**工具与 Skill 的分工**：工具回答“能做什么”（查课程、检索文档、算学分），Skill 回答“该怎么做”"
     "（先调哪个工具、要不要追问、用什么关键词检索、输出成什么格式、何时拒答）。调整回答策略只需修改 Skill "
     "描述文件，**不必改代码、不必重新打包插件**。", size=9.5)

h("5.2 插件提供的能力与集成", 2)
bullet("**工具与命令注册**：课程、教师、方案、条款查询，文档语义检索，学分核算与冲突查询；并提供 /选课问答、/课程、/教师、/学分等快捷入口，降低新生使用门槛。", size=9.5)
bullet("**配置与密钥**：数据库、向量库、模型 API、数据版本、调用限额统一在插件配置中管理；密钥走配置或环境变量，不硬编码。", size=9.5)
bullet("**界面卡片**：课程卡、教师卡、来源引用、学分进度视图、追问建议，让结构化信息比纯文本更易读。", size=9.5)
bullet("**会话画像**：首启动引导选择专业与年级，可补充已修课程，此后所有回答自动带入画像。", size=9.5)

h("5.3 部署演进", 2)
flow("单机轻量（本地数据库 + 本地向量索引，用于个人验证与快速试点）→ "
     "学院服务器（数据集中维护、多人共用）→ 校内服务化（对接统一身份认证与教务数据，合规要求最高）")

# ================================================================ 六
h("六、实施计划与风险", 1)

h("6.1 分阶段实施", 2)
table(
    ["阶段", "周期", "主要任务", "交付物"],
    [
        ["一、需求与数据摸底", "1～2 周", "确定专业与年级范围、盘点数据与授权、收集高频问题", "需求与数据清单、高频问题 Top 50"],
        ["二、数据治理与库表落地", "2～3 周", "解析清洗、结构化抽取与人工校验、建别名表、数据入库", "结构化数据库、术语与别名表"],
        ["三、RAG 与 Skill 体系", "3～4 周", "打通检索链路、建设 Skill 集合、建评测集并做回归优化", "完整检索链路、Skill 集合、评测报告"],
        ["四、插件封装与试点", "2～3 周", "插件结构搭建、工具与命令注册、UI 卡片、单专业试点", "可安装插件包、试点评估报告"],
    ],
    widths=[3.2, 1.6, 7.4, 4.4],
)
caption("表 5　分阶段实施计划（合计约 8～12 周）")

h("6.2 主要风险与应对", 2)
table(
    ["风险", "应对措施"],
    [
        ["数据不规范、多版本并存", "建立版本主数据；关键字段 100% 人工复核；按年级强制过滤检索范围"],
        ["幻觉与错误引用", "引用程序化校验 + 数字一致性复核 + 低置信度拒答 + 每周人工抽检闭环"],
        ["教师相关内容引发争议", "只用公开信息、只做客观呈现、不做评价与排名，并预设话术与审核机制"],
        ["政策或数据变动导致过期", "版本化与失效标记；定期更新；回答中标注数据版本与截止时间"],
        ["模型 API 成本与限流", "模型分级路由；高频问答缓存；设置调用限额与降级策略"],
        ["个人数据安全与合规", "最小化采集；身份与范围校验；脱敏日志；必要的授权与审批"],
    ],
    widths=[4.4, 12.2],
)
caption("表 6　主要风险与应对")

h("6.3 验收指标", 2)
table(
    ["答案准确率", "引用可溯源率", "高频问题覆盖率", "拒答合理性", "幻觉率"],
    [["≥ 90%", "100%", "≥ 85%（Top 50）", "≥ 95%", "≤ 2%"]],
    widths=[3.32] * 5,
)
caption("表 7　试点期验收指标")

# ================================================================ 七
h("七、待确认事项", 1)
numbered("**目标范围**：面向哪些专业、哪些年级？是否预留扩展到其他专业的接口。", size=9.5)
numbered("**数据授权**：培养方案、教学大纲、教师数据、开课计划的获取渠道与使用授权。", size=9.5)
numbered("**教务对接**：是否需要实时读取个人学业数据？接口形式与权限审批流程如何。", size=9.5)
numbered("**模型与预算**：用公有云模型 API 还是校内自建服务？预算、配额，以及数据是否允许出校。", size=9.5)
numbered("**教师口径**：教师信息展示范围如何界定？是否需要学院先行审核。", size=9.5)
numbered("**交付与维护**：插件在何处运行、学生如何获取；数据更新与 Skill 维护由谁负责。", size=9.5)
numbered("**合规审查**：是否需要通过学校信息化或数据安全部门审查。", size=9.5)

# ================================================================ 八
h("八、下一步工作建议", 1)
numbered("**先做数据，后做模型**：尽快拿到一份完整、带版本的培养方案与课程库，这决定了项目约 70% 的效果上限。", size=9.5)
numbered("**先做窄场景，后做全场景**：第一版只做“培养方案解读 + 课程事实查询”，把引用溯源与拒答机制打磨扎实，再扩展教师画像与规划建议。", size=9.5)
numbered("**先建评测集，后调优**：动手开发前先整理 50 条真实高频问题及其标准答案与应命中证据，作为后续所有优化的判据。", size=9.5)
numbered("**合规先行**：教师数据与个人学业数据相关功能，建议在开发前完成口径确认与必要审批。", size=9.5)
para("（本文为精简摘要。完整方案见同目录《计算机类新生选课与培养方案问答助手_初步设想.docx》，共 22 页，含架构、库表、"
     "Skill 设计、场景示例等详细内容。）", size=8.5, space_after=0, color=GREY)

# ================================================================ 保存
out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "计算机类新生选课与培养方案问答助手_方案摘要_4页版.docx")

cp = doc.core_properties
cp.title = "计算机类新生选课与培养方案问答助手 —— 方案摘要（精简版）"
cp.subject = "基于大模型与检索增强生成（RAG）· 4 页精简版"
cp.category = "方案文档"
cp.comments = "精简版摘要，供快速阅读与汇报使用；完整方案见《初步设想》文档。"
cp.keywords = "RAG; 大模型; DSH 插件; Skill; 培养方案; 选课"

doc.save(out_path)

pages = EST["cm"] / USABLE_H
print("saved:", out_path)
print("size:", os.path.getsize(out_path), "bytes")
print("estimated content height: %.1f cm -> %.2f pages (usable %.1f cm/page)"
      % (EST["cm"], pages, USABLE_H))
